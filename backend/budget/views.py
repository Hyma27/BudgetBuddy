from openai import models
from django.db import models
from django.db.models import Sum
from django.db.models.functions import TruncMonth
from datetime import datetime
from openpyxl import Workbook
from io import BytesIO
from django.http import HttpResponse

from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework import status
from rest_framework_simplejwt.tokens import RefreshToken
from rest_framework.permissions import IsAuthenticated

from .models import User,Income,Expense,Budget,SavingsGoal,Notification,Report
from .serializers import UserRegistrationSerializer,IncomeSerializer,ExpenseSerializer,BudgetSerializer,SavingsGoalSerializer,NotificationSerializer,ReportSerializer


class RegisterView(APIView):

    def post(self, request):
        serializer = UserRegistrationSerializer(data=request.data)

        if serializer.is_valid():
            serializer.save()

            return Response(
                {
                    'message': 'User registered successfully'
                },
                status=status.HTTP_201_CREATED
            )

        return Response(
            serializer.errors,
            status=status.HTTP_400_BAD_REQUEST
        )


class LoginView(APIView):

    def post(self, request):
        email = request.data.get('email')
        password = request.data.get('password')

        try:
            user = User.objects.get(email=email)
        except User.DoesNotExist:
            return Response(
                {'error': 'Invalid email or password'},
                status=status.HTTP_401_UNAUTHORIZED
            )

        if not user.check_password(password):
            return Response(
                {'error': 'Invalid email or password'},
                status=status.HTTP_401_UNAUTHORIZED
            )

        refresh = RefreshToken.for_user(user)

        return Response({
            'refresh': str(refresh),
            'access': str(refresh.access_token),
            'role': user.role,
        })
        
        
class ProfileView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        return Response({
            'message': 'You are authenticated successfully',
            'username': request.user.username,
            'email': request.user.email,
            'role': request.user.role,
        })        
        
class IncomeView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request, income_id=None):

    # If no specific income ID is given,
    # return all incomes of the logged-in user
        if income_id is None:
            incomes = Income.objects.filter(
              user=request.user
            ).order_by('-income_date', '-id')

            serializer = IncomeSerializer(
              incomes,
              many=True
            )

            return Response(serializer.data)
        
        # If a specific income ID is given,
        # return only that income if it belongs to the logged-in user
        try:
            income = Income.objects.get(
               id=income_id,
               user=request.user
            )
        except Income.DoesNotExist:
            return Response(
               {
                   "error": "Income not found"
               },
               status=status.HTTP_404_NOT_FOUND
           )

        serializer = IncomeSerializer(income)

        return Response(
            serializer.data,
            status=status.HTTP_200_OK
        )

    def post(self, request):
        serializer = IncomeSerializer(
            data=request.data
        )

        if serializer.is_valid():
            serializer.save(user=request.user)

            return Response(
                serializer.data,
                status=status.HTTP_201_CREATED
            )

        return Response(
            serializer.errors,
            status=status.HTTP_400_BAD_REQUEST
        )
     
    def put(self, request, income_id):
        try:
            income = Income.objects.get(
                id=income_id,
                user=request.user
            )
        except Income.DoesNotExist:
            return Response(
                {
                    "error": "Income not found"
                },
                status=status.HTTP_404_NOT_FOUND
            )

        serializer = IncomeSerializer(
            income,
            data=request.data
        )

        if serializer.is_valid():
            serializer.save()

            return Response(
                serializer.data,
                status=status.HTTP_200_OK
            )

        return Response(
            serializer.errors,
            status=status.HTTP_400_BAD_REQUEST
        )

    def delete(self, request, income_id):
        try:
            income = Income.objects.get(
                id=income_id,
                user=request.user
            )
        except Income.DoesNotExist:
            return Response(
                {
                    "error": "Income not found"
                },
                status=status.HTTP_404_NOT_FOUND
            )

        income.delete()

        return Response(
            {
                "message": "Income deleted successfully"
            },
            status=status.HTTP_200_OK
        )   
class ExpenseView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        expenses = Expense.objects.filter(
            user=request.user
        ).order_by('-expense_date', '-id')

        serializer = ExpenseSerializer(
            expenses,
            many=True
        )

        return Response(serializer.data)

    def post(self, request):
        serializer = ExpenseSerializer(
            data=request.data
        )

        if serializer.is_valid():
            expense = serializer.save(user=request.user)

            # Check matching monthly budget
            budget = Budget.objects.filter(
                user=request.user,
                category=expense.category,
                period="Monthly"
            ).first()

            if budget:

                # Calculate spending for the same category
                # and same month as the new expense
                total_spending = Expense.objects.filter(
                    user=request.user,
                    category=budget.category,
                    expense_date__year=expense.expense_date.year,
                    expense_date__month=expense.expense_date.month
                ).aggregate(
                    total=models.Sum('amount')
                )['total'] or 0

                # Calculate budget utilization percentage
                utilization = (
                    total_spending / budget.amount
                ) * 100

                notification_type = None
                message = None

                # Budget exceeded/reached
                if utilization >= 100:
                   notification_type = "Budget Limit Alert"
                   message = (
                        f"Your {budget.category} monthly budget of "
                        f"₹{budget.amount} has been reached or exceeded. "
                        f"Current spending: ₹{total_spending} "
                        f"({utilization:.2f}%) "
                        f"for {expense.expense_date.year}-"
                        f"{expense.expense_date.month:02d}."
                    )

        # Budget near limit
                elif utilization >= 80:
                    notification_type = "Budget Near Limit Alert"

                    message = (
                        f"Your {budget.category} monthly budget is near its limit. "
                        f"Budget: ₹{budget.amount}, "
                        f"Current spending: ₹{total_spending} "
                        f"({utilization:.2f}%) "
                        f"for {expense.expense_date.year}-"
                        f"{expense.expense_date.month:02d}."
                    )
            

               # Create notification only when threshold is reached
                if notification_type:

                        month_label = (
                            f"{expense.expense_date.year}-"
                            f"{expense.expense_date.month:02d}"
                        )

                        existing_notification = (
                            Notification.objects.filter(
                                user=request.user,
                                notification_type=notification_type,
                            )
                            .filter(message__contains=budget.category)
                            .filter(message__contains=month_label)
                            .exists()
                        )

                        if not existing_notification:
                            Notification.objects.create(
                                user=request.user,
                                notification_type=notification_type,
                                message=message
                            )
                        # Create notification only when threshold is reached
                        if notification_type:

                                    existing_notification = Notification.objects.filter(
                                        user=request.user,
                                        notification_type=notification_type,
                                        message__contains=budget.category,
                                        is_read=False
                                    ).filter(
                                        message__contains=(
                                            f"{expense.expense_date.year}-"
                                            f"{expense.expense_date.month:02d}"
                                    )
                                    ).exists()
                                        
                    

                                    if not existing_notification:
                                        Notification.objects.create(
                                            user=request.user,
                                            notification_type=notification_type,
                                            message=message
                                        )

            return Response(
                serializer.data,
                status=status.HTTP_201_CREATED
            )

        return Response(
            serializer.errors,
            status=status.HTTP_400_BAD_REQUEST
        )

    def put(self, request, expense_id):
        try:
            expense = Expense.objects.get(
                id=expense_id,
                user=request.user
            )
        except Expense.DoesNotExist:
            return Response(
                {
                    "error": "Expense not found"
                },
                status=status.HTTP_404_NOT_FOUND
            )

        serializer = ExpenseSerializer(
            expense,
            data=request.data
        )

        if serializer.is_valid():
            serializer.save()

            return Response(
                serializer.data,
                status=status.HTTP_200_OK
            )

        return Response(
            serializer.errors,
            status=status.HTTP_400_BAD_REQUEST
        )

    def delete(self, request, expense_id):
        try:
            expense = Expense.objects.get(
                id=expense_id,
                user=request.user
            )
        except Expense.DoesNotExist:
            return Response(
                {
                    "error": "Expense not found"
                },
                status=status.HTTP_404_NOT_FOUND
            )

        expense.delete()

        return Response(
            {
                "message": "Expense deleted successfully"
            },
            status=status.HTTP_200_OK
        )
        
        
class BudgetView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        budgets = Budget.objects.filter(
            user=request.user
        ).order_by('-id')

        serializer = BudgetSerializer(
            budgets,
            many=True
        )

        return Response(serializer.data)

    def post(self, request):
        serializer = BudgetSerializer(
            data=request.data
        )

        if serializer.is_valid():
            serializer.save(user=request.user)

            return Response(
                serializer.data,
                status=status.HTTP_201_CREATED
            )

        return Response(
            serializer.errors,
            status=status.HTTP_400_BAD_REQUEST
        )
    
    
class SavingsGoalView(APIView):
    permission_classes = [IsAuthenticated]

    # Get user's savings goals
    def get(self, request):
        savings_goals = SavingsGoal.objects.filter(
            user=request.user
        )

        serializer = SavingsGoalSerializer(
            savings_goals,
            many=True
        )

        return Response(serializer.data)

    # Create a savings goal
    def post(self, request):
        serializer = SavingsGoalSerializer(
            data=request.data
        )

        if serializer.is_valid():
            savings_goal = serializer.save(
                user=request.user
            )

            return Response(
                SavingsGoalSerializer(savings_goal).data,
                status=status.HTTP_201_CREATED
            )

        return Response(
            serializer.errors,
            status=status.HTTP_400_BAD_REQUEST
       )

    # Update a savings goal
    def put(self, request, goal_id=None):

        try:
            savings_goal = SavingsGoal.objects.get(
                id=goal_id,
                user=request.user
            )

        except SavingsGoal.DoesNotExist:
            return Response(
                {"error": "Savings goal not found"},
                status=status.HTTP_404_NOT_FOUND
            )

        serializer = SavingsGoalSerializer(
            savings_goal,
            data=request.data,
            partial=True
        )

        if serializer.is_valid():

            savings_goal = serializer.save()

            # Check savings goal milestone
            if (
                savings_goal.current_amount
                >= savings_goal.target_amount
            ):

                # Prevent duplicate unread milestone notifications
                existing_notification = Notification.objects.filter(
                    user=request.user,
                    notification_type="Savings Goal Milestone",
                    message__contains=savings_goal.goal_name,
                    is_read=False
                ).exists()

                if not existing_notification:

                    Notification.objects.create(
                        user=request.user,
                        notification_type="Savings Goal Milestone",
                        message=(
                            f"Congratulations! You have reached your "
                            f"savings goal '{savings_goal.goal_name}' "
                            f"of ₹{savings_goal.target_amount}."
                        )
                    )

            return Response(
                SavingsGoalSerializer(savings_goal).data,
                status=status.HTTP_200_OK
            )

        return Response(
            serializer.errors,
            status=status.HTTP_400_BAD_REQUEST
        )
    
    
class NotificationView(APIView):
    permission_classes = [IsAuthenticated]

    # Get all notifications
    def get(self, request):
        unread_only = request.query_params.get("unread")
        notifications = Notification.objects.filter(user=request.user).order_by('-created_at')
        if unread_only == "true":
             notifications = notifications.filter(is_read=False)

        notifications = notifications.order_by('-created_at')
        
        serializer = NotificationSerializer(notifications, many=True)

        return Response(serializer.data)

    # Mark notification as read
    def put(self, request, notification_id):
        try:
            notification = Notification.objects.get(
                id=notification_id,
                user=request.user
            )
        except Notification.DoesNotExist:
            return Response(
                {"error": "Notification not found"},
                status=status.HTTP_404_NOT_FOUND
            )

        notification.is_read = True
        notification.save()

        serializer = NotificationSerializer(notification)

        return Response(serializer.data)
    

class ReportView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        reports = Report.objects.filter(
            user=request.user
        ).order_by('-generated_at')

        report_data = []

        for report in reports:

            # Get report month
            try:
                selected_month = datetime.strptime(
                    report.period,
                    "%Y-%m"
                )
            except ValueError:
                continue

            # Monthly income
            incomes = Income.objects.filter(
                user=request.user,
                income_date__year=selected_month.year,
                income_date__month=selected_month.month
            )

            # Monthly expenses
            expenses = Expense.objects.filter(
                user=request.user,
                expense_date__year=selected_month.year,
                expense_date__month=selected_month.month
            )

            # Calculate totals
            total_income = incomes.aggregate(
                total=Sum("amount")
            )["total"] or 0

            total_expenses = expenses.aggregate(
                total=Sum("amount")
            )["total"] or 0

            savings = total_income - total_expenses

            # Get user's budgets
            budgets = Budget.objects.filter(
                user=request.user
            )

            budget_data = [
                {
                    "category": budget.category,
                    "amount": budget.amount,
                    "period": budget.period
                }
                for budget in budgets
            ]

            report_data.append({
                "id": report.id,
                "report_type": report.report_type,
                "period": report.period,
                "generated_at": report.generated_at,
                "file_reference": report.file_reference,

                "summary": {
                    "total_income": total_income,
                    "total_expenses": total_expenses,
                    "savings": savings
                },

                "budgets": budget_data,

                "expense_count": expenses.count(),
                "income_count": incomes.count()
            })

        return Response(
            report_data,
            status=status.HTTP_200_OK
        )

    def post(self, request):

        # Get requested month
        month = request.data.get("month")

        if not month:
            return Response(
                {"error": "Month is required. Use YYYY-MM."},
                status=status.HTTP_400_BAD_REQUEST
            )

        # Validate month
        try:
            selected_month = datetime.strptime(
                month,
                "%Y-%m"
            )
        except ValueError:
            return Response(
                {"error": "Invalid month format. Use YYYY-MM."},
                status=status.HTTP_400_BAD_REQUEST
            )

        # Get user's monthly income
        incomes = Income.objects.filter(
            user=request.user,
            income_date__year=selected_month.year,
            income_date__month=selected_month.month
        )

        # Get user's monthly expenses
        expenses = Expense.objects.filter(
            user=request.user,
            expense_date__year=selected_month.year,
            expense_date__month=selected_month.month
        )

        # Calculate totals
        total_income = incomes.aggregate(
            total=Sum("amount")
        )["total"] or 0

        total_expenses = expenses.aggregate(
            total=Sum("amount")
        )["total"] or 0

        savings = total_income - total_expenses

        # Get user's budgets
        budgets = Budget.objects.filter(
            user=request.user
        )

        # Create report record
        report = Report.objects.create(
            user=request.user,
            report_type="Monthly Financial Report",
            period=month,
            file_reference=f"BudgetBuddy_{month}.xlsx"
        )

        # Create notification
        Notification.objects.create(
            user=request.user,
            notification_type="Monthly Report Generated",
            message=(
                f"Your monthly financial report for {month} "
                f"has been generated successfully."
            )
        )

        # Return JSON response
        return Response(
            {
                "message": "Monthly report generated successfully.",
                "report": {
                    "id": report.id,
                    "report_type": report.report_type,
                    "period": report.period,
                    "generated_at": report.generated_at,
                    "file_reference": report.file_reference,

                    "summary": {
                        "total_income": total_income,
                        "total_expenses": total_expenses,
                        "savings": savings
                    },

                    "expense_count": expenses.count(),
                    "income_count": incomes.count(),
                    "budget_count": budgets.count()
                }
            },
            status=status.HTTP_201_CREATED
        )

    def excel(self, request, report_id):

        # Get only the logged-in user's report
        try:
            report = Report.objects.get(
                id=report_id,
                user=request.user
            )
        except Report.DoesNotExist:
            return Response(
                {"error": "Report not found"},
                status=status.HTTP_404_NOT_FOUND
            )

        # Validate report month
        try:
            selected_month = datetime.strptime(
                report.period,
                "%Y-%m"
            )
        except ValueError:
            return Response(
                {"error": "Invalid report period"},
                status=status.HTTP_400_BAD_REQUEST
            )

        # Get user's monthly income
        incomes = Income.objects.filter(
            user=request.user,
            income_date__year=selected_month.year,
            income_date__month=selected_month.month
        )

        # Get user's monthly expenses
        expenses = Expense.objects.filter(
            user=request.user,
            expense_date__year=selected_month.year,
            expense_date__month=selected_month.month
        )

        # Get user's budgets
        budgets = Budget.objects.filter(
            user=request.user
        )

        # Calculate totals
        total_income = incomes.aggregate(
            total=Sum("amount")
        )["total"] or 0

        total_expenses = expenses.aggregate(
            total=Sum("amount")
        )["total"] or 0

        savings = total_income - total_expenses

        # Create Excel workbook
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Monthly Report"

        # Report information
        worksheet["A1"] = "BudgetBuddy - Monthly Financial Report"
        worksheet["A2"] = f"User: {request.user.username}"
        worksheet["A3"] = f"Period: {report.period}"

        # Financial summary
        worksheet["A5"] = "Financial Summary"

        worksheet["A6"] = "Total Income"
        worksheet["B6"] = float(total_income)

        worksheet["A7"] = "Total Expenses"
        worksheet["B7"] = float(total_expenses)

        worksheet["A8"] = "Savings"
        worksheet["B8"] = float(savings)

        # Expenses
        worksheet["A10"] = "Expenses"

        worksheet["A11"] = "Title"
        worksheet["B11"] = "Amount"
        worksheet["C11"] = "Category"
        worksheet["D11"] = "Date"
        worksheet["E11"] = "Description"

        row = 12

        for expense in expenses.order_by(
            "expense_date",
            "id"
        ):
            worksheet.cell(
                row=row,
                column=1,
                value=expense.title
            )

            worksheet.cell(
                row=row,
                column=2,
                value=float(expense.amount)
            )

            worksheet.cell(
                row=row,
                column=3,
                value=expense.category
            )

            worksheet.cell(
                row=row,
                column=4,
                value=str(expense.expense_date)
            )

            worksheet.cell(
                row=row,
                column=5,
                value=expense.description or ""
            )

            row += 1

        # Income
        row += 2

        worksheet.cell(
            row=row,
            column=1,
            value="Income"
        )

        row += 1

        worksheet.cell(
            row=row,
            column=1,
            value="Source"
        )

        worksheet.cell(
            row=row,
            column=2,
            value="Amount"
        )

        worksheet.cell(
            row=row,
            column=3,
            value="Date"
        )

        worksheet.cell(
            row=row,
            column=4,
            value="Description"
        )

        row += 1

        for income in incomes.order_by(
            "income_date",
            "id"
        ):
            worksheet.cell(
                row=row,
                column=1,
                value=income.source
            )

            worksheet.cell(
                row=row,
                column=2,
                value=float(income.amount)
            )

            worksheet.cell(
                row=row,
                column=3,
                value=str(income.income_date)
            )

            worksheet.cell(
                row=row,
                column=4,
                value=income.description or ""
            )

            row += 1

        # Budgets
        row += 2

        worksheet.cell(
            row=row,
            column=1,
            value="Budgets"
        )

        row += 1

        worksheet.cell(
            row=row,
            column=1,
            value="Category"
        )

        worksheet.cell(
            row=row,
            column=2,
            value="Amount"
        )

        worksheet.cell(
            row=row,
            column=3,
            value="Period"
        )

        row += 1

        for budget in budgets:
            worksheet.cell(
                row=row,
                column=1,
                value=budget.category
            )

            worksheet.cell(
                row=row,
                column=2,
                value=float(budget.amount)
            )

            worksheet.cell(
                row=row,
                column=3,
                value=budget.period
            )

            row += 1

        # Column widths
        worksheet.column_dimensions["A"].width = 30
        worksheet.column_dimensions["B"].width = 18
        worksheet.column_dimensions["C"].width = 20
        worksheet.column_dimensions["D"].width = 18
        worksheet.column_dimensions["E"].width = 35

        # Save workbook in memory
        excel_file = BytesIO()

        workbook.save(excel_file)

        excel_file.seek(0)

        # Return Excel file
        response = HttpResponse(
            excel_file.getvalue(),
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            )
        )
        

        response["Content-Disposition"] = (
            f'attachment; filename="BudgetBuddy_{report.period}.xlsx"'
        )

        return response

class ReportExcelView(ReportView):
    def get(self, request, report_id):
        return self.excel(request, report_id)     
        
class ProtectedView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):
        return Response({
            "message": f"Hello {request.user.username}, you are authenticated!"
        })
        
        
class AnalyticsView(APIView):
    permission_classes = [IsAuthenticated]

    def get(self, request):

        # Logged-in user's data only
        incomes = Income.objects.filter(
            user=request.user
        )

        expenses = Expense.objects.filter(
            user=request.user
        )

        # -----------------------------
        # Date / Month Filtering
        # -----------------------------

        start_date = request.query_params.get("start_date")
        end_date = request.query_params.get("end_date")
        month = request.query_params.get("month")

        # Month filter: YYYY-MM
        if month:
            try:
                selected_month = datetime.strptime(
                    month,
                    "%Y-%m"
                )

                incomes = incomes.filter(
                    income_date__year=selected_month.year,
                    income_date__month=selected_month.month
                )

                expenses = expenses.filter(
                    expense_date__year=selected_month.year,
                    expense_date__month=selected_month.month
                )

            except ValueError:
                return Response(
                    {
                        "error": "Invalid month format. Use YYYY-MM."
                    },
                    status=status.HTTP_400_BAD_REQUEST
                )

        # Date range filter
        elif start_date or end_date:

            if start_date:
                try:
                    start = datetime.strptime(
                        start_date,
                        "%Y-%m-%d"
                    ).date()

                except ValueError:
                    return Response(
                        {
                            "error": "Invalid start_date. Use YYYY-MM-DD."
                        },
                        status=status.HTTP_400_BAD_REQUEST
                    )

                incomes = incomes.filter(
                    income_date__gte=start
                )

                expenses = expenses.filter(
                    expense_date__gte=start
                )

            if end_date:
                try:
                    end = datetime.strptime(
                        end_date,
                        "%Y-%m-%d"
                    ).date()

                except ValueError:
                    return Response(
                        {
                            "error": "Invalid end_date. Use YYYY-MM-DD."
                        },
                        status=status.HTTP_400_BAD_REQUEST
                    )

                incomes = incomes.filter(
                    income_date__lte=end
                )

                expenses = expenses.filter(
                    expense_date__lte=end
                )

        # -----------------------------
        # Total Income
        # -----------------------------

        total_income = incomes.aggregate(
            total=Sum("amount")
        )["total"] or 0

        # -----------------------------
        # Total Expenses
        # -----------------------------

        total_expenses = expenses.aggregate(
            total=Sum("amount")
        )["total"] or 0

        # -----------------------------
        # Savings
        # -----------------------------

        savings = total_income - total_expenses

        # -----------------------------
        # Category-wise Expenses
        # -----------------------------

        category_data = expenses.values(
            "category"
        ).annotate(
            total=Sum("amount")
        ).order_by("category")

        category_summary = [
            {
                "category": item["category"],
                "total": item["total"]
            }
            for item in category_data
        ]

        # -----------------------------
        # Monthly Trends
        # -----------------------------

        income_monthly = incomes.annotate(
            month=TruncMonth("income_date")
        ).values(
            "month"
        ).annotate(
            total=Sum("amount")
        ).order_by("month")

        expense_monthly = expenses.annotate(
            month=TruncMonth("expense_date")
        ).values(
            "month"
        ).annotate(
            total=Sum("amount")
        ).order_by("month")

        income_trends = [
            {
                "month": item["month"].strftime("%Y-%m"),
                "total": item["total"]
            }
            for item in income_monthly
        ]

        expense_trends = [
            {
                "month": item["month"].strftime("%Y-%m"),
                "total": item["total"]
            }
            for item in expense_monthly
        ]

        # -----------------------------
        # Final Response
        # -----------------------------

        return Response(
            {
                "summary": {
                    "total_income": total_income,
                    "total_expenses": total_expenses,
                    "savings": savings
                },

                "category_wise_expenses": category_summary,

                "monthly_trends": {
                    "income": income_trends,
                    "expenses": expense_trends
                }
            },
        )
        
    